"""
Прогноз остатков и заявка поставщику.

Считает по каждой позиции скорость продаж (шт/день) из `sales_orders` — той
самой истории завершённых заказов, которую наполняет
`backend/services/sales_sync_worker.py`, — делит на неё текущий остаток из
каталога и получает «дней до нуля». Новых таблиц не заводит: всё живое,
считается на лету по тем же данным, что уже есть.

Чего здесь сознательно НЕТ и почему:

* **Истории остатков.** Мы знаем остаток только «сейчас» (`_load_products`),
  а не то, каким он был неделю назад. Значит, дни, когда товар физически
  лежал в нуле и продаваться не мог, неотличимы от дней, когда он просто не
  продавался, — скорость на таких SKU занижена. Честнее занизить, чем
  придумать: заказ по заниженной скорости приведёт к недозаказу, который
  видно на следующей неделе, а по завышенной — к затовариванию, которое
  видно через квартал.
* **Закупочных цен.** Их никто в системе не хранит, поэтому «сумма заявки»
  считается по ЦЕНЕ ПРОДАЖИ и подписана в выгрузке именно так. Выдавать её
  за себестоимость нельзя.

**Мастер-группы остатка — отдельная строка, а не N строк товаров.** Если
несколько артикулов делят один физический остаток
(`master_stock_groups.shared_stock`, см. `stock_sync_worker.py`), то их нельзя
прогнозировать по отдельности: у каждого в каталоге Kaspi лежит одно и то же
число, и сумма «дней до нуля» по пяти карточкам одной группы завысила бы
запас впятеро. Такие артикулы схлопываются в одну строку: остаток — общий,
скорость — сумма скоростей участников, умноженная на размер комплекта
(`deduct_per_unit`, то же правило, по которому списывает воркер).

День берётся АЛМАТИНСКИЙ (UTC+5, см. ALMATY_UTC_OFFSET_HOURS в
routers/analytics.py): продавец думает в своих сутках. Текущий, ещё не
закончившийся день из окна исключён — иначе каждое утро скорость всего
каталога проваливалась бы, потому что «сегодня» продано почти ничего.
"""

import datetime as dt
import io
import json
import logging
import math
from typing import Optional

from sqlalchemy import select

from backend.product_name_overrides import is_placeholder_name, resolve_product_name
from backend.state_helpers import (
    _get_sync_db_session,
    _load_master_stock_groups,
    _load_products,
)

logger = logging.getLogger(__name__)

ALMATY_UTC_OFFSET_HOURS = 5

# Дефолты, которые может переопределить продавец в запросе.
DEFAULT_WINDOW_DAYS = 60       # глубина истории для «медленной» скорости
DEFAULT_RECENT_DAYS = 14       # окно «быстрой» скорости (свежий тренд)
DEFAULT_LEAD_TIME_DAYS = 14    # сколько идёт поставка от поставщика
DEFAULT_COVER_DAYS = 14        # на сколько дней продаж закупаемся сверх поставки
DEFAULT_HORIZON_DAYS = 7       # «закончится в течение недели»
DEFAULT_SLOW_DAYS = 90         # запаса больше чем на столько дней — «продаётся слишком медленно»

# Вес свежего окна в итоговой скорости. Спрос на Kaspi живёт по кривой
# «буст в выдаче → пик → спад» (то же наблюдение, что в analytics.py::abc-xyz),
# поэтому средняя за два месяца систематически врёт на товарах в фазе роста
# и в фазе затухания. Берём смесь с перевесом в сторону последних двух недель.
_RECENT_WEIGHT = 0.6

BUCKETS = ("out_of_stock", "critical", "warning", "ok", "slow", "dead")


def _local_day(created_at: dt.datetime) -> dt.date:
    return (created_at + dt.timedelta(hours=ALMATY_UTC_OFFSET_HOURS)).date()


def _norm(value) -> str:
    return str(value or "").strip()


def _velocity(sold: int, first_day: dt.date, window_start: dt.date, window_end: dt.date) -> float:
    """Штук в день на активном отрезке окна.

    Знаменатель — не вся длина окна, а дни от ПЕРВОЙ продажи (или от начала
    окна, если товар продавался и раньше) до конца окна. Иначе карточка,
    выставленная три дня назад и продавшая 6 штук, получила бы 0.1 шт/день
    вместо 2 шт/день, и её бы никогда не показало в «скоро кончится».
    """
    start = max(window_start, first_day)
    if start > window_end:
        return 0.0
    days = (window_end - start).days + 1
    return sold / days if days > 0 else 0.0


def _blend(recent: float, long: float, first_day: Optional[dt.date], recent_start: dt.date) -> float:
    """Итоговая скорость. Если товар начал продаваться уже внутри свежего
    окна, «медленная» скорость посчитана по тем же самым дням — смешивать
    её с самой собой смысла нет, берём свежую как есть."""
    if first_day is None:
        return 0.0
    if first_day >= recent_start:
        return recent
    return _RECENT_WEIGHT * recent + (1.0 - _RECENT_WEIGHT) * long


def _bucket(stock: int, velocity: float, sold: int, days_to_zero: Optional[float],
            horizon_days: int, lead_time_days: int, slow_days: int) -> str:
    if velocity <= 0:
        # Ноль продаж за всё окно при живом остатке — это замороженные деньги,
        # а не «всё хорошо»: отдельная категория, ради которой продавец и
        # открывает такой отчёт вторым делом после «что кончается».
        return "dead" if stock > 0 and sold == 0 else "ok"
    if stock <= 0:
        return "out_of_stock"
    if days_to_zero is not None and days_to_zero <= horizon_days:
        return "critical"
    if days_to_zero is not None and days_to_zero <= lead_time_days:
        # Успеет кончиться раньше, чем приедет поставка, даже если заказать
        # прямо сейчас — но позже, чем через неделю.
        return "warning"
    if days_to_zero is not None and days_to_zero >= slow_days:
        # Продажи есть, но остатка хватит на квартал и дальше: деньги лежат
        # в товаре, который оборачивается слишком медленно. От "dead" это
        # отличается тем, что товар живой — его имеет смысл разгонять
        # (цена, реклама), а не только списывать. Закупать такое не надо:
        # recommended_qty у него всё равно 0, потому что целевой запас
        # (срок поставки + покрытие) заведомо меньше текущего остатка.
        return "slow"
    return "ok"


def _select_items_for_response(items: list, max_items: Optional[int]) -> tuple:
    """Урезает items для API-ответа, не теряя ничего важного.

    Дашборд опрашивает /stock-forecast каждые 5 минут с каждого открытого
    клиента. На живых данных Q-STEPPE (5155 позиций, из них 4423 «мёртвых»
    без единой продажи) полный список весит ~3 МБ JSON на один такой опрос —
    тот же класс проблемы, что уже ловили на /api/stock-groups (HANDOFF §3.52,
    p95=27.6с из-за лишних данных в ответе).

    Правило: сначала ВСЁ «горящее» (out_of_stock/critical/warning) и ВСЁ, что
    реально попадает в заявку (recommended_qty > 0) — это никогда не режется,
    иначе таблица могла бы молча не показать позицию, которую summary уже
    посчитал в reorder_units.

    Остаток бюджета делится между спокойными вёдрами (ok/slow/dead) поровну,
    а не отдаётся первым по сортировке: сортировка идёт по возрастанию «дней
    до нуля», поэтому сплошной срез съел бы бюджет на «ok» и оставил бы
    фильтры «продаётся медленно» и «не продаётся» пустыми ровно на тех
    каталогах, где они нужнее всего. Внутри этих двух вёдер вперёд идут
    позиции с наибольшей замороженной суммой — с них продавец и начнёт.

    Экспорт в Excel (build_purchase_order_xlsx) получает ПОЛНЫЙ список
    отдельным вызовом без max_items — урезание касается только экрана.
    """
    if max_items is None or len(items) <= max_items:
        return items, False
    priority, rest = [], []
    for it in items:
        if it["bucket"] in ("out_of_stock", "critical", "warning") or it["recommended_qty"] > 0:
            priority.append(it)
        else:
            rest.append(it)
    budget = max(0, max_items - len(priority))
    order = {id(it): n for n, it in enumerate(rest)}

    by_bucket: dict = {}
    for it in rest:
        by_bucket.setdefault(it["bucket"], []).append(it)
    for bucket in ("slow", "dead"):
        if bucket in by_bucket:
            by_bucket[bucket].sort(key=lambda it: -(it["stock"] * (it.get("price") or 0)))

    share = budget // len(by_bucket) if by_bucket else 0
    picked, spare = [], []
    for rows in by_bucket.values():
        picked.extend(rows[:share])
        spare.extend(rows[share:])
    # Ведро могло не выбрать свою квоту (в нём просто меньше строк) — остаток
    # бюджета отдаём обратно в общий котёл, в исходном порядке сортировки.
    spare.sort(key=lambda it: order[id(it)])
    picked.extend(spare[: max(0, budget - len(picked))])
    picked.sort(key=lambda it: order[id(it)])
    return priority + picked, True


def build_forecast(
    products: list,
    groups: dict,
    daily_by_sku: dict,
    today: dt.date,
    window_days: int = DEFAULT_WINDOW_DAYS,
    recent_days: int = DEFAULT_RECENT_DAYS,
    lead_time_days: int = DEFAULT_LEAD_TIME_DAYS,
    cover_days: int = DEFAULT_COVER_DAYS,
    horizon_days: int = DEFAULT_HORIZON_DAYS,
    slow_days: int = DEFAULT_SLOW_DAYS,
    max_items: Optional[int] = None,
) -> dict:
    """Чистое ядро: никаких походов в БД, только уже загруженные данные.

    `daily_by_sku` — {merchant_sku: {date: проданное количество}}.
    `products` — каталог в формате `_load_products`.
    `groups` — мастер-группы в формате `_load_master_stock_groups`.
    `max_items` — см. `_select_items_for_response`; None (по умолчанию) —
    список не режется, summary в любом случае считается по полному набору.
    """
    # Сегодняшний день ещё не закончился — считаем по закрытым суткам.
    window_end = today - dt.timedelta(days=1)
    window_start = window_end - dt.timedelta(days=window_days - 1)
    recent_start = max(window_start, window_end - dt.timedelta(days=recent_days - 1))

    catalog_by_msku: dict = {}
    for p in products:
        m_sku = _norm(p.get("merchantSku"))
        if m_sku:
            catalog_by_msku.setdefault(m_sku, p)

    # ── Скорость по каждому артикулу ────────────────────────────────────
    per_sku: dict = {}
    for m_sku, daily in daily_by_sku.items():
        in_window = {d: q for d, q in daily.items() if window_start <= d <= window_end and q > 0}
        if not in_window:
            per_sku[m_sku] = {"sold": 0, "sold_recent": 0, "velocity": 0.0,
                              "velocity_recent": 0.0, "velocity_long": 0.0,
                              "first_day": None, "last_day": None}
            continue
        first_day = min(in_window)
        last_day = max(in_window)
        sold = sum(in_window.values())
        sold_recent = sum(q for d, q in in_window.items() if d >= recent_start)
        v_long = _velocity(sold, first_day, window_start, window_end)
        v_recent = _velocity(sold_recent, first_day, recent_start, window_end)
        per_sku[m_sku] = {
            "sold": sold,
            "sold_recent": sold_recent,
            "velocity": _blend(v_recent, v_long, first_day, recent_start),
            "velocity_recent": v_recent,
            "velocity_long": v_long,
            "first_day": first_day,
            "last_day": last_day,
        }

    # ── Раскладка артикулов по мастер-группам ───────────────────────────
    # Группа с выключенным is_active остаток больше не транслирует — её
    # участники прогнозируются как обычные товары, каждый по своему остатку.
    # Индексированный матчинг обязателен: прямой перебор стоит
    # O(артикулов x групп) и на Бархане (1589 x 286) держал один прогон
    # прогноза 14 секунд — ручку, которую дашборд опрашивает у каждого
    # открытого клиента. С индексом — доли секунды.
    from backend.services.stock_sync_worker import build_group_index, match_group_for_sku_indexed

    active_groups = {gid: g for gid, g in (groups or {}).items() if g.get("is_active")}
    group_members: dict = {}      # group_id -> [merchant_sku]
    sku_group: dict = {}          # merchant_sku -> group_id
    if active_groups:
        index = build_group_index(active_groups)
        for m_sku in set(per_sku) | set(catalog_by_msku):
            gid, _g = match_group_for_sku_indexed(m_sku, index)
            if gid:
                sku_group[m_sku] = gid
                group_members.setdefault(gid, []).append(m_sku)

    items = []

    def _finish(row: dict) -> dict:
        velocity = row["velocity"]
        stock = row["stock"]
        days_to_zero = (stock / velocity) if velocity > 0 else None
        row["days_to_zero"] = round(days_to_zero, 1) if days_to_zero is not None else None
        # Дата «закончится» есть только у того, что ещё не закончилось: при
        # нулевом остатке days_to_zero == 0 и сюда попадала сегодняшняя дата,
        # то есть «закончится сегодня» у товара, который стоит в нуле уже
        # неделю. Пусто честнее.
        row["runs_out_on"] = (
            (today + dt.timedelta(days=int(days_to_zero))).isoformat()
            if days_to_zero is not None and 0 < days_to_zero < 3650 else None
        )
        row["bucket"] = _bucket(stock, velocity, row["sold"], days_to_zero,
                                horizon_days, lead_time_days, slow_days)
        # Целевой запас — на срок поставки плюс период, который она должна
        # закрыть. Заказываем разницу с тем, что уже лежит на складе.
        if velocity > 0:
            target = velocity * (lead_time_days + cover_days)
            row["recommended_qty"] = max(0, int(math.ceil(target - stock)))
        else:
            row["recommended_qty"] = 0
        price = row.get("price") or 0
        row["reorder_cost"] = round(row["recommended_qty"] * price, 2)
        row["velocity"] = round(velocity, 3)
        row["velocity_recent"] = round(row["velocity_recent"], 3)
        row["velocity_long"] = round(row["velocity_long"], 3)
        return row

    # Строки мастер-групп
    for gid, members in group_members.items():
        g = active_groups[gid]
        per_unit = max(1, int(g.get("deduct_per_unit", 1) or 1))
        stats = [per_sku.get(m, {}) for m in members]
        sold_members = sorted(m for m in members if per_sku.get(m, {}).get("sold"))
        items.append(_finish({
            "kind": "group",
            "key": gid,
            "name": g.get("group_name") or gid,
            "merchant_sku": ", ".join(sold_members) or None,
            "kaspi_sku": None,
            "external_code": g.get("external_code"),
            "image": next((catalog_by_msku[m].get("image") for m in members if m in catalog_by_msku), None),
            "status": "Активен",
            "in_catalog": True,
            "members_count": len(members),
            "deduct_per_unit": per_unit,
            "stock": int(g.get("shared_stock", 0) or 0),
            "sold": sum(s.get("sold", 0) for s in stats),
            "sold_recent": sum(s.get("sold_recent", 0) for s in stats),
            "velocity": sum(s.get("velocity", 0.0) for s in stats) * per_unit,
            "velocity_recent": sum(s.get("velocity_recent", 0.0) for s in stats) * per_unit,
            "velocity_long": sum(s.get("velocity_long", 0.0) for s in stats) * per_unit,
            "price": _group_price(members, catalog_by_msku),
        }))

    # Строки обычных товаров
    unmatched_skus = 0
    unmatched_sold = 0
    for m_sku in sorted(set(per_sku) | set(catalog_by_msku)):
        if m_sku in sku_group:
            continue
        stats = per_sku.get(m_sku, {})
        p = catalog_by_msku.get(m_sku)
        sold = stats.get("sold", 0)
        # Артикула нет в текущем каталоге. На живых данных Бархана таких
        # оказалось 53 штуки, и это НЕ потерянные продажи: там старые
        # написания артикулов с лишними нулями ("X147-00006" при живом
        # "X147-0006"), под которыми товар продавался до переименования.
        # Складского остатка у такой строки нет by design, поэтому она
        # неизбежно попадала в «кончился» и раздувала главную цифру отчёта
        # больше чем в полтора раза. Заказать по несуществующему артикулу
        # всё равно нельзя — держим их за пределами вёдер, но считаем и
        # показываем количество, чтобы расхождение не пропало молча.
        if p is None:
            if sold:
                unmatched_skus += 1
                unmatched_sold += sold
            continue
        # Архивная карточка без единой продажи в окне — это снятый с продажи
        # хлам, которого в каталоге активного магазина бывают сотни. К закупу
        # он отношения не имеет. А вот архивная С продажами — это как раз
        # «кончилось и сняли с витрины», самая ценная строка отчёта.
        if sold == 0 and p.get("status") != "Активен":
            continue
        name = p.get("name") or m_sku
        if is_placeholder_name(name, m_sku):
            name = resolve_product_name(m_sku, None) or name
        items.append(_finish({
            "kind": "sku",
            "key": m_sku,
            "name": name,
            "merchant_sku": m_sku,
            "kaspi_sku": _norm(p.get("sku")),
            "external_code": None,
            "image": p.get("image"),
            "status": p.get("status") or "",
            "in_catalog": True,
            "members_count": 1,
            "deduct_per_unit": 1,
            "stock": int(p.get("stock") or 0),
            "sold": sold,
            "sold_recent": stats.get("sold_recent", 0),
            "velocity": stats.get("velocity", 0.0),
            "velocity_recent": stats.get("velocity_recent", 0.0),
            "velocity_long": stats.get("velocity_long", 0.0),
            "price": int(p.get("price_num") or 0),
        }))

    # Самое срочное — наверх. Позиции без скорости (dead/ok) уезжают в конец.
    items.sort(key=lambda it: (
        it["days_to_zero"] is None,
        it["days_to_zero"] if it["days_to_zero"] is not None else 0,
        -it["sold"],
    ))

    # Summary — всегда по ПОЛНОМУ набору, до урезания items ниже. Иначе на
    # большом каталоге сумма по видимым строкам разошлась бы с заголовочной
    # цифрой (Q-STEPPE: reorder_units посчитан по 357 позициям, а таблица
    # показала бы только часть из них).
    summary = {b: 0 for b in BUCKETS}
    for it in items:
        summary[it["bucket"]] = summary.get(it["bucket"], 0) + 1
    reorder = [it for it in items if it["recommended_qty"] > 0]

    response_items, truncated = _select_items_for_response(items, max_items)

    return {
        "generated_at": dt.datetime.utcnow().isoformat(),
        "params": {
            "window_days": window_days,
            "recent_days": recent_days,
            "lead_time_days": lead_time_days,
            "cover_days": cover_days,
            "horizon_days": horizon_days,
            "slow_days": slow_days,
        },
        "window": {"from": window_start.isoformat(), "to": window_end.isoformat()},
        "summary": {
            **summary,
            "total_items": len(items),
            "reorder_positions": len(reorder),
            "reorder_units": sum(it["recommended_qty"] for it in reorder),
            "reorder_cost": round(sum(it["reorder_cost"] for it in reorder), 2),
            "frozen_cost": round(sum(it["stock"] * (it.get("price") or 0)
                                     for it in items if it["bucket"] == "dead"), 2),
            # Те же замороженные деньги, но в живом, просто медленном товаре —
            # считаем отдельно от "dead", иначе одна цифра смешивала бы то,
            # что надо распродавать, с тем, что надо разгонять.
            "slow_cost": round(sum(it["stock"] * (it.get("price") or 0)
                                   for it in items if it["bucket"] == "slow"), 2),
            # Артикулы из истории продаж, которых нет в текущем каталоге —
            # в вёдра не входят, см. комментарий у отбора строк.
            "unmatched_skus": unmatched_skus,
            "unmatched_sold": unmatched_sold,
        },
        # items_truncated=True значит: на экране не все "ok"/"dead" позиции
        # без закупа (все actionable — здесь всегда, полностью). Экспорт в
        # Excel бьёт по полному списку отдельным вызовом compute_stock_forecast
        # без max_items.
        "items_truncated": truncated,
        "items": response_items,
    }


def _group_price(members: list, catalog_by_msku: dict) -> int:
    """Цена мастер-группы для оценки суммы заявки — максимум по участникам.
    Своей цены у группы нет, а участники это, как правило, одна и та же вещь
    в разных фасовках; минимум занизил бы стоимость запаса."""
    prices = [int(catalog_by_msku[m].get("price_num") or 0) for m in members if m in catalog_by_msku]
    return max(prices) if prices else 0


# ── Загрузка данных ─────────────────────────────────────────────────────

def load_daily_sales(store_id: str, since: dt.date) -> dict:
    """{merchant_sku: {алматинский день: продано штук}} из sales_orders.

    Ключ — merchantSku (поле "sku" в entries_payload), а не числовой
    productCode: у части заказов Kaspi отдаёт productCode нераскодированным
    (пустая строка), и группировка по нему разваливает один товар на два —
    ровно тот баг, что уже был пойман вживую в sales_export.py::_entry_group_key.
    """
    daily: dict = {}
    session = _get_sync_db_session()
    if not session:
        return daily
    try:
        from backend.db.models import SalesOrderDB

        # Запас в сутки с каждой стороны — created_at в БД лежит в UTC, а
        # режем мы по алматинскому дню (UTC+5).
        since_dt = dt.datetime.combine(since - dt.timedelta(days=1), dt.time.min)
        rows = session.scalars(
            select(SalesOrderDB).where(
                SalesOrderDB.store_id == store_id,
                SalesOrderDB.created_at >= since_dt,
            )
        ).all()
        for r in rows:
            try:
                entries = json.loads(r.entries_payload)
            except Exception:
                continue
            day = _local_day(r.created_at)
            for e in entries:
                m_sku = _norm(e.get("sku"))
                if not m_sku:
                    continue
                qty = int(e.get("quantity") or 0)
                if qty <= 0:
                    continue
                daily.setdefault(m_sku, {})
                daily[m_sku][day] = daily[m_sku].get(day, 0) + qty
    finally:
        session.close()
    return daily


def compute_stock_forecast(
    store_id: str,
    window_days: int = DEFAULT_WINDOW_DAYS,
    recent_days: int = DEFAULT_RECENT_DAYS,
    lead_time_days: int = DEFAULT_LEAD_TIME_DAYS,
    cover_days: int = DEFAULT_COVER_DAYS,
    horizon_days: int = DEFAULT_HORIZON_DAYS,
    slow_days: int = DEFAULT_SLOW_DAYS,
    max_items: Optional[int] = None,
) -> dict:
    """Прогноз для магазина: тянет каталог, мастер-группы и историю продаж.

    `max_items` — см. `_select_items_for_response`. Ручка экрана вызывает с
    лимитом (большой каталог не должен весить мегабайты на каждый 5-минутный
    опрос дашборда), выгрузка в Excel — без лимита.
    """
    today = (dt.datetime.utcnow() + dt.timedelta(hours=ALMATY_UTC_OFFSET_HOURS)).date()
    since = today - dt.timedelta(days=window_days)
    return build_forecast(
        products=_load_products(store_id),
        groups=_load_master_stock_groups(store_id=store_id),
        daily_by_sku=load_daily_sales(store_id, since),
        today=today,
        window_days=window_days,
        recent_days=recent_days,
        lead_time_days=lead_time_days,
        cover_days=cover_days,
        horizon_days=horizon_days,
        slow_days=slow_days,
        max_items=max_items,
    )


# ── Выгрузка заявки поставщику ──────────────────────────────────────────

PURCHASE_HEADERS = [
    "Артикул",
    "Код поставщика",
    "Наименование",
    "Тип",
    "Остаток, шт.",
    "Продано за период, шт.",
    "Скорость, шт./день",
    "Дней до нуля",
    "Закончится",
    "ЗАКАЗАТЬ, шт.",
    "Цена продажи, ₸",
    "Сумма по цене продажи, ₸",
]


def build_purchase_order_xlsx(forecast: dict, only_reorder: bool = True) -> bytes:
    """Заявка поставщику в Excel по готовому прогнозу.

    Колонка «Сумма» считается по ЦЕНЕ ПРОДАЖИ, а не по закупочной — закупочных
    цен система не знает (их негде взять из Kaspi), и подписана колонка именно
    так, чтобы её не приняли за себестоимость заявки.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    rows = [it for it in forecast["items"] if not only_reorder or it["recommended_qty"] > 0]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заявка"

    params = forecast["params"]
    ws.append([
        f"Заявка поставщику · окно продаж {forecast['window']['from']} — {forecast['window']['to']}"
        f" · срок поставки {params['lead_time_days']} дн. · запас на {params['cover_days']} дн."
    ])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])

    ws.append(PURCHASE_HEADERS)
    header_row = ws.max_row
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    urgent = PatternFill("solid", fgColor="FFC7CE")
    for it in rows:
        ws.append([
            it.get("merchant_sku") or "",
            it.get("external_code") or "",
            it.get("name") or "",
            "Мастер-группа" if it["kind"] == "group" else "Товар",
            it["stock"],
            it["sold"],
            it["velocity"],
            it["days_to_zero"] if it["days_to_zero"] is not None else "",
            it["runs_out_on"] or "",
            it["recommended_qty"],
            it.get("price") or 0,
            it["reorder_cost"],
        ])
        if it["bucket"] in ("out_of_stock", "critical"):
            for cell in ws[ws.max_row]:
                cell.fill = urgent

    total_row = ws.max_row + 1
    ws.cell(row=total_row, column=3, value="ИТОГО").font = Font(bold=True)
    ws.cell(row=total_row, column=10, value=sum(it["recommended_qty"] for it in rows)).font = Font(bold=True)
    ws.cell(row=total_row, column=12, value=round(sum(it["reorder_cost"] for it in rows), 2)).font = Font(bold=True)

    widths = [18, 16, 46, 15, 12, 18, 15, 13, 13, 14, 16, 20]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
